"""Tests de l'export multi-jours (app/history_export.py)."""
import os

import pytest

from app import config, db
from app import history_export


def _seed(conn, make_df, rayon, jour, rows, statut="ok"):
    df = make_df(rows)
    imp = db.create_import(conn, rayon, jour, f"f-{jour}.xlsx", f"h-{rayon}-{jour}", statut,
                           nb_articles=len(df))
    db.insert_snapshot(conn, imp, rayon, jour, df)
    return imp


ROWS_J1 = [
    {"Code": 1, "Libellé": "OEUFS FRAIS X12", "Stock": 10, "Px vente": 1500,
     "Px revient": 1000, "Marge %": 33.33, "Couv. ": 5},
    {"Code": 2, "Libellé": "BOEUF HACHE 500G", "Stock": -3, "Px vente": 2000,
     "Px revient": 2500, "Marge %": -25.0, "Couv. ": 999},
    {"Code": 3, "Libellé": "LAIT ENTIER 1L", "Stock": 50, "Px vente": 900,
     "Px revient": 700, "Marge %": 22.22, "Couv. ": None},
]
ROWS_J2 = [
    {"Code": 1, "Libellé": "OEUFS FRAIS X12", "Stock": 4, "Px vente": 1500,
     "Px revient": 1000, "Marge %": 33.33, "Couv. ": 3},
    {"Code": 2, "Libellé": "BOEUF HACHE 500G", "Stock": -8, "Px vente": 2000,
     "Px revient": 2500, "Marge %": -25.0, "Couv. ": 999},
    {"Code": 4, "Libellé": "EGG NOODLES 250G", "Stock": 20, "Px vente": 1200,
     "Px revient": 800, "Marge %": 33.33, "Couv. ": 7},
]


@pytest.fixture()
def seeded(fresh_db, make_df, tmp_path, monkeypatch):
    monkeypatch.setattr(history_export, "EXPORTS_DIR", str(tmp_path / "exports"))
    monkeypatch.setattr(history_export, "PUBLIC_BASE", "https://x.test/exports")
    with db.lock_conn() as conn:
        _seed(conn, make_df, "frais-surgele", "2026-09-10", ROWS_J1)
        _seed(conn, make_df, "frais-surgele", "2026-09-11", ROWS_J2)
    return tmp_path


def test_export_tous_les_jours(seeded):
    res = history_export.export_history("frais-surgele", "OEUF,EGG", format="xlsx")
    assert res["success"] is True
    assert res["jours_utilises"] == ["2026-09-10", "2026-09-11"]
    assert res["nb_articles"] == 3  # codes 1, 2, 4
    assert res["nb_lignes"] == 5
    assert res["partiel"] is False
    assert os.path.getsize(res["fichier"]) > 0


def test_export_plage_dates(seeded):
    res = history_export.export_history("frais-surgele", "OEUF", date_debut="2026-09-11",
                                        date_fin="2026-09-11", format="csv")
    assert res["success"] is True
    assert res["jours_utilises"] == ["2026-09-11"]
    assert res["nb_lignes"] == 2  # codes 1 (OEUFS) et 2 (BOEUF, pas d'exclusion)
    assert res["fichier"].endswith(".csv")


def test_filtre_inclut_boeuf(seeded):
    res = history_export.export_history("frais-surgele", "OEUF", format="csv")
    import pandas as pd
    df = pd.read_csv(res["fichier"], dtype=str)
    assert any("BOEUF" in (v or "") for v in df["Libelle"])


def test_valeurs_particulieres_conservees(seeded):
    res = history_export.export_history("frais-surgele", "OEUF,EGG", format="xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(res["fichier"], data_only=True)
    assert "Historique" in wb.sheetnames and "Resume" in wb.sheetnames
    ws = wb["Historique"]
    assert [c.value for c in ws[1]] == ["Jour", "Code article", "Libelle", "Stock",
                                       "Prix vente FDJ", "Prix revient FDJ", "Marge %", "Couverture"]
    lignes = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(lignes) == res["nb_lignes"] == 5
    boeuf = [l for l in lignes if l[2] and "BOEUF" in l[2]]
    assert boeuf and all(l[3] < 0 for l in boeuf)  # stocks negatifs conserves
    assert all(l[6] < 0 for l in boeuf)  # marges negatives conservees
    assert all(l[7] == 999 for l in boeuf)  # sentinelle 999 conservee
    assert all(isinstance(l[1], str) for l in lignes)  # codes en texte


def test_valeurs_nulles_et_ordre(seeded):
    res = history_export.export_history("frais-surgele", "LAIT", format="csv")
    import pandas as pd
    df = pd.read_csv(res["fichier"], dtype=str)
    assert len(df) == 1
    assert df.iloc[0]["Couverture"] in (None, "", "nan") or pd.isna(df.iloc[0]["Couverture"])
    res2 = history_export.export_history("frais-surgele", "OEUF,EGG", format="csv")
    df2 = pd.read_csv(res2["fichier"], dtype=str)
    assert list(zip(df2["Jour"], df2["Code article"])) == sorted(zip(df2["Jour"], df2["Code article"]))


def test_sans_import(fresh_db, tmp_path, monkeypatch):
    monkeypatch.setattr(history_export, "EXPORTS_DIR", str(tmp_path / "exports"))
    res = history_export.export_history("frais-surgele", "OEUF")
    assert res["success"] is False


def test_rayon_inconnu_et_format_invalide(seeded):
    assert history_export.export_history("nope", "OEUF")["success"] is False
    assert history_export.export_history("frais-surgele", "OEUF", format="pdf")["success"] is False
    assert history_export.export_history("frais-surgele", "OEUF", date_debut="11-09-2026")["success"] is False
    assert history_export.export_history("frais-surgele", "")["success"] is False
