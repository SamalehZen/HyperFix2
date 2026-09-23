"""Tests de l'Excel intelligent : pivot, selections, tris, couleurs, memoire."""
import json

import pandas as pd
import pytest

from app import config, db
from app import excel_intelligent as xi


def _seed(conn):
    """3 articles x 3 jours : 44774-like (chute), stable, hausse."""
    jours = ["2026-09-14", "2026-09-15", "2026-09-16"]
    data = {
        44774: ("ANGURIA MELON 80ML", [31.93, 31.93, 10.06], [100, 100, 20]),
        10001: ("ARTICLE STABLE", [20.0, 20.0, 20.0], [10, 10, 10]),
        10002: ("ARTICLE HAUSSE", [5.0, 6.0, 9.0], [200, 200, 200]),
    }
    for jour in jours:
        rows = []
        for code, (lib, marges, stocks) in data.items():
            rows.append({"Code": code, "Libellé": lib, "Marge %": marges[jours.index(jour)],
                         "Stock": stocks[jours.index(jour)], "Px vente": 100.0, "Px revient": 80.0,
                         "Fournisseur": "F", "Couv. ": 5.0, "Valeur stock   PRMP": stocks[jours.index(jour)] * 80.0})
        df = pd.DataFrame(rows)
        for c in config.REQUIRED_COLUMNS:
            if c not in df.columns:
                df[c] = None
        iid = db.create_import(conn, "frais-surgele", jour, f"f-{jour}.xlsx", f"h-{jour}", "ok", nb_articles=3)
        db.insert_snapshot(conn, iid, "frais-surgele", jour, df)


@pytest.fixture()
def seeded(fresh_db, monkeypatch):
    monkeypatch.setattr(config, "rayon_ids", lambda: ["frais-surgele", "epicerie-salee"])
    with fresh_db.lock_conn() as conn:
        _seed(conn)
    return fresh_db


def _plan(**kw):
    p = {"rayon": "frais-surgele", "indicateur": "marge", "selection": "baisses",
         "seuil_baisse_pts": 5, "double_classement": True, "couleurs": True, "resume": True}
    p.update(kw)
    return p


def test_pivot_baisses_et_tris(seeded, tmp_path, monkeypatch):
    monkeypatch.setattr(xi, "EXPORTS_DIR", str(tmp_path))
    monkeypatch.setattr(xi, "PUBLIC_BASE", "https://x/exports")
    monkeypatch.setattr(xi, "LAST_PLAN_FILE", str(tmp_path / "last.json"))
    res = xi.build_excel(_plan())
    assert res["success"], res.get("erreur")
    assert res["nb_articles"] == 1  # seul 44774 baisse >= 5 pts
    assert res["nb_jours"] == 3
    assert all(v["ok"] for v in res["verifications"]), res["verifications"]


def test_double_classement_et_couleurs(seeded, tmp_path, monkeypatch):
    monkeypatch.setattr(xi, "EXPORTS_DIR", str(tmp_path))
    monkeypatch.setattr(xi, "PUBLIC_BASE", "https://x/exports")
    monkeypatch.setattr(xi, "LAST_PLAN_FILE", str(tmp_path / "last.json"))
    res = xi.build_excel(_plan(selection="tous"))
    assert res["success"], res.get("erreur")
    from openpyxl import load_workbook
    wb = load_workbook(res["fichier"])
    assert set(wb.sheetnames) == {"Prios argent", "Prix à corriger", "Résumé"}
    ws = wb["Prios argent"]
    codes = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert len(codes) == len(set(codes)) == 3  # 1 ligne par article
    entetes = [ws.cell(row=1, column=c).value for c in range(9, 12)]
    assert entetes == ["Marge 14/09", "Marge 15/09", "Marge 16/09"]
    fixes = [ws.cell(row=1, column=c).value for c in range(1, 9)]
    assert fixes == ["Code article", "Libellé", "Fournisseur", "Stock",
                     "Px revient", "Px vente", "Promo active", "Cause"]
    # 44774 : chute 16/09 -> cellule rouge
    ligne = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == 44774)
    assert str(ws.cell(row=ligne, column=11).fill.start_color.rgb).endswith("FFC7CE")
    # Hausse 10002 le 16/09 -> verte
    ligne2 = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == 10002)
    assert str(ws.cell(row=ligne2, column=11).fill.start_color.rgb).endswith("C6EFCE")


def test_selection_vide_honnete(seeded, tmp_path, monkeypatch):
    monkeypatch.setattr(xi, "EXPORTS_DIR", str(tmp_path))
    monkeypatch.setattr(xi, "LAST_PLAN_FILE", str(tmp_path / "last.json"))
    res = xi.build_excel(_plan(selection="negatifs"))
    assert not res["success"]
    assert "Aucun article" in res["erreur"]


def test_memoire_dernier_plan(seeded, tmp_path, monkeypatch):
    monkeypatch.setattr(xi, "EXPORTS_DIR", str(tmp_path))
    monkeypatch.setattr(xi, "PUBLIC_BASE", "https://x/exports")
    monkeypatch.setattr(xi, "LAST_PLAN_FILE", str(tmp_path / "last.json"))
    r1 = xi.build_excel(_plan())
    assert r1["success"]
    # "Refais avec le stock" : seul l'indicateur change, le reste est conserve.
    r2 = xi.build_excel({"indicateur": "stock"}, base="dernier")
    assert r2["success"], r2.get("erreur")
    assert r2["plan"]["indicateur"] == "stock"
    assert r2["plan"]["selection"] == "baisses"
    assert r2["plan"]["rayon"] == "frais-surgele"


def test_rayon_inconnu_refuse(tmp_path, monkeypatch):
    monkeypatch.setattr(xi, "LAST_PLAN_FILE", str(tmp_path / "last.json"))
    res = xi.build_excel({"rayon": "nope"})
    assert not res["success"]
    assert "Rayon inconnu" in res["erreur"]


def test_verificateur_detecte_doublon(seeded, tmp_path, monkeypatch):
    monkeypatch.setattr(xi, "EXPORTS_DIR", str(tmp_path))
    monkeypatch.setattr(xi, "PUBLIC_BASE", "https://x/exports")
    monkeypatch.setattr(xi, "LAST_PLAN_FILE", str(tmp_path / "last.json"))
    res = xi.build_excel(_plan(selection="tous", double_classement=False))
    assert res["success"]
    from openpyxl import load_workbook
    wb = load_workbook(res["fichier"])
    ws = wb["Analyse"]
    # Sabotage : duplique la premiere ligne de donnees.
    ws.append([ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)])
    wb.save(res["fichier"])
    ctx = {"jours": res["jours_utilises"], "label_ind": "Marge",
           "titres_attendus": ["Analyse"],
           "ordres": {"Analyse": [44774, 10001, 10002]},
           "rouges": {}, "points": [], "feuille_point": "Analyse",
           "selection": "tous", "resume_demande": True}
    checks = xi.verifier_classeur(res["fichier"], ctx)
    assert any((not v["ok"]) and "uniques" in v["nom"] for v in checks)


def test_resolve_rayon_tolerant(seeded):
    assert config.resolve_rayon("frais-surgele") == "frais-surgele"
    assert config.resolve_rayon("Frais surgelé") == "frais-surgele"
    assert config.resolve_rayon("FRAIS-SURGELE") == "frais-surgele"
    assert config.resolve_rayon("  frais-surgele  ") == "frais-surgele"
    assert config.resolve_rayon("Épicerie salée") == "epicerie-salee"
    assert config.resolve_rayon("nope") is None
    assert config.resolve_rayon("") is None
    assert config.resolve_rayon(None) is None


def test_plan_from_args_facon_luna():
    # Args exacts envoyes par Luna le 2026-09-17 (sans plan_json).
    plan = xi.plan_from_args({"type": "marge", "rayon": "Frais surgelé",
                              "format": "xlsx", "resume": True,
                              "inclure_stock": True})
    assert plan["rayon"] == "Frais surgelé"  # resolu plus tard en id canonique
    assert plan["indicateur"] == "marge"
    assert plan.get("resume") is True
    assert "format" not in plan and "inclure_stock" not in plan
    plan2 = xi.plan_from_args({"rayon_id": "frais-surgele", "type_export": "stock",
                               "pivot": "marge", "mots_cles": "OEUF,EGG"})
    assert plan2["rayon"] == "frais-surgele"
    assert plan2["indicateur"] == "stock"
    assert plan2["selection"] == {"mots": "OEUF,EGG"}


def test_appel_plat_luna_aboutit(seeded, tmp_path, monkeypatch):
    # Bout en bout : les args plats de Luna produisent un fichier verifie.
    monkeypatch.setattr(xi, "EXPORTS_DIR", str(tmp_path))
    monkeypatch.setattr(xi, "PUBLIC_BASE", "https://x/exports")
    monkeypatch.setattr(xi, "LAST_PLAN_FILE", str(tmp_path / "last.json"))
    plan = {"rayon": "frais-surgele"}
    plan.update(xi.plan_from_args({"type": "marge", "rayon": "Frais surgelé",
                                   "format": "xlsx", "resume": True}))
    plan["rayon"] = config.resolve_rayon(plan["rayon"])
    res = xi.build_excel(plan)
    assert res["success"], res.get("erreur")
    assert res["nb_articles"] == 1


# ------------------------------------------------- jeu n2 : priorites ----
def _seed2(conn):
    """Priorites (gros codes en tete), diagnostic, dormants, ruptures, fantome."""
    jours = ["2026-09-14", "2026-09-15", "2026-09-16"]
    arts = {
        999001: dict(lib="TOP PERTE", marges=[50.0, 50.0, 10.0], stocks=[1000]*3,
                      vente=200.0, revient=100.0, fourn="F1", couv=10.0, valeur=100000.0),
        999002: dict(lib="TOP CHUTE", marges=[90.0, 90.0, 5.0], stocks=[1]*3,
                      vente=100.0, revient=50.0, fourn="F2", couv=50.0, valeur=50.0),
        100: dict(lib="PETIT", marges=[30.0, 30.0, 28.0], stocks=[5]*3,
                   vente=100.0, revient=70.0, fourn="F3", couv=10.0, valeur=350.0),
        200: dict(lib="DORMANT", marges=[10.0]*3, stocks=[100]*3,
                   vente=100.0, revient=90.0, fourn="F4", couv=999.0, valeur=50000.0),
        300: dict(lib="RUPTURE", marges=[15.0]*3, stocks=[5, 3, 2],
                   vente=100.0, revient=80.0, fourn="F5", couv=5.0, valeur=200.0),
        400: dict(lib="FANTOME", marges=[None]*3, stocks=[10]*3,
                   vente=None, revient=None, fourn="F6", couv=10.0, valeur=500.0),
        500: dict(lib="STRUCT", marges=[10.0, 10.0, -5.0], stocks=[4]*3,
                   vente=60.0, revient=80.0, fourn="F7", couv=10.0, valeur=320.0),
        600: dict(lib="INCOH", marges=[20.0, 20.0, -5.0], stocks=[6]*3,
                   vente=100.0, revient=80.0, fourn="F8", couv=10.0, valeur=480.0),
        700: dict(lib="NEGSTOCK", marges=[40.0, 40.0, 10.0], stocks=[-100]*3,
                   vente=200.0, revient=100.0, fourn="F9", couv=10.0, valeur=0.0),
    }
    for jour in jours:
        i = jours.index(jour)
        rows = [{"Code": c, "Libellé": a["lib"], "Fournisseur": a["fourn"],
                 "Marge %": a["marges"][i], "Stock": a["stocks"][i],
                 "Px vente": a["vente"], "Px revient": a["revient"],
                 "Couv. ": a["couv"], "Valeur stock   PRMP": a["valeur"]}
                for c, a in arts.items()]
        df = pd.DataFrame(rows)
        for col in config.REQUIRED_COLUMNS:
            if col not in df.columns:
                df[col] = None
        iid = db.create_import(conn, "frais-surgele", jour, f"g-{jour}.xlsx",
                               f"h2-{jour}", "ok", nb_articles=len(rows))
        db.insert_snapshot(conn, iid, "frais-surgele", jour, df)


@pytest.fixture()
def seeded2(fresh_db, monkeypatch):
    monkeypatch.setattr(config, "rayon_ids", lambda: ["frais-surgele", "epicerie-salee"])
    with fresh_db.lock_conn() as conn:
        _seed2(conn)
    return fresh_db


def _plan2(**kw):
    p = {"rayon": "frais-surgele", "indicateur": "marge", "selection": "tous",
         "double_classement": True, "couleurs": True, "resume": True}
    p.update(kw)
    return p


def _run2(tmp_path, monkeypatch, plan):
    monkeypatch.setattr(xi, "EXPORTS_DIR", str(tmp_path))
    monkeypatch.setattr(xi, "PUBLIC_BASE", "https://x/exports")
    monkeypatch.setattr(xi, "LAST_PLAN_FILE", str(tmp_path / "last.json"))
    return xi.build_excel(plan)


def test_priorite_avant_coupe(seeded2, tmp_path, monkeypatch):
    # Plafond a 2 : les tops MONDIAUX survivent quel que soit leur code.
    monkeypatch.setattr(xi, "MAX_ARTICLES", 2)
    res = _run2(tmp_path, monkeypatch, _plan2())
    assert res["success"], res.get("erreur")
    assert res["partiel"] is True
    from openpyxl import load_workbook
    wb = load_workbook(res["fichier"])
    prios = [wb["Prios argent"].cell(row=r, column=1).value for r in (2, 3)]
    assert prios == [999001, 600], prios  # top pertes : 40000 puis 90
    prix = [wb["Prix à corriger"].cell(row=r, column=1).value for r in (2, 3)]
    assert prix == [400, 999002], prix  # fantome en tete, puis top chute


def test_cause_diagnostic(seeded2, tmp_path, monkeypatch):
    res = _run2(tmp_path, monkeypatch, _plan2())
    assert res["success"], res.get("erreur")
    assert res["causes"] == {"STRUCTURELLE": 1, "INCOHERENTE": 1}, res["causes"]
    from openpyxl import load_workbook
    wb = load_workbook(res["fichier"])
    ws = wb["Prix à corriger"]
    causes = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=8).value
              for r in range(2, ws.max_row + 1)}
    assert causes[500] == "STRUCTURELLE", causes
    assert causes[600] == "INCOHERENTE", causes
    assert causes[400] == "DONNEE", causes


def test_dormants_ruptures(seeded2, tmp_path, monkeypatch):
    res = _run2(tmp_path, monkeypatch, _plan2())
    assert res["success"], res.get("erreur")
    assert res["dormants"] == {"nb": 1, "capital": 50000}, res["dormants"]
    assert res["ruptures_nb"] == 1
    assert res["fantomes_nb"] == 1
    assert res["feuilles"]["Dormants"] == {"lignes": 1, "total": 1}
    assert res["feuilles"]["A commander"] == {"lignes": 1, "total": 1}
    from openpyxl import load_workbook
    wb = load_workbook(res["fichier"])
    assert wb["Dormants"].cell(row=2, column=1).value == 200
    assert wb["Dormants"].cell(row=2, column=5).value == 50000
    assert wb["A commander"].cell(row=2, column=1).value == 300
    assert set(wb.sheetnames) == {"Prios argent", "Prix à corriger", "Dormants",
                                  "A commander", "Résumé"}


def test_stock_negatif_sans_capital(seeded2, tmp_path, monkeypatch):
    # 700 : chute -30 mais stock -100 -> perte 0 -> fin de Prios (pas d'argent
    # en jeu), mais present dans Prix (la chute reste a corriger).
    res = _run2(tmp_path, monkeypatch, _plan2())
    assert res["success"], res.get("erreur")
    from openpyxl import load_workbook
    wb = load_workbook(res["fichier"])
    prios = [wb["Prios argent"].cell(row=r, column=1).value
             for r in range(2, wb["Prios argent"].max_row + 1)]
    assert prios[0] == 999001 and prios[-1] == 700, prios
    prix = [wb["Prix à corriger"].cell(row=r, column=1).value
            for r in range(2, wb["Prix à corriger"].max_row + 1)]
    assert 700 in prix and prix[0] == 400, prix


def test_appel_plat_luna_aboutit_bis(seeded, tmp_path, monkeypatch):
    # Doublon volontaire du test ci-dessus (meme nom historique) : bout en
    # bout avec les args plats de Luna.
    monkeypatch.setattr(xi, "EXPORTS_DIR", str(tmp_path))
    monkeypatch.setattr(xi, "PUBLIC_BASE", "https://x/exports")
    monkeypatch.setattr(xi, "LAST_PLAN_FILE", str(tmp_path / "last.json"))
    plan = {"rayon": "frais-surgele"}
    plan.update(xi.plan_from_args({"type": "marge", "rayon": "Frais surgelé",
                                   "format": "xlsx", "resume": True}))
    plan["rayon"] = config.resolve_rayon(plan["rayon"])
    res = xi.build_excel(plan)
    assert res["success"], res.get("erreur")
    assert res["nb_articles"] == 1


# --------------------------------- jeu n3 : marge<0 ET stock>0 meme jour ----
def _seed_negpos(conn):
    """Marge<0 ET stock>0 le MEME jour (+ pieges + valeurs texte)."""
    jours = ["2026-07-30", "2026-07-31", "2026-08-01"]
    data = {
        900001: ("CAS POSITIF", [-5.0, 3.0, 3.0], [10, 10, 10]),
        900005: ("CAS POSITIF MULTI", [-2.0, -9.0, 4.0], [3, 7, 7]),
        900006: ("CAS TEXTE", ["-7.5", None, "abc"], ["4", "5", "6"]),
        900002: ("PIEGE JOURS SEPARES", [-8.0, 5.0, 5.0], [0, 12, 12]),
        900003: ("PIEGE STOCK NUL/NEG", [-3.0, -4.0, 6.0], [0, -5, 9]),
        900004: ("PIEGE MARGE POSITIVE", [6.0, 6.0, 6.0], [20, 20, 20]),
    }
    for jour in jours:
        i = jours.index(jour)
        rows = [{"Code": c, "Libellé": lib, "Marge %": marges[i], "Stock": stocks[i],
                 "Px vente": 100.0, "Px revient": 80.0, "Fournisseur": "F",
                 "Couv. ": 5.0, "Valeur stock   PRMP": 80.0}
                for c, (lib, marges, stocks) in data.items()]
        df = pd.DataFrame(rows)
        for col in config.REQUIRED_COLUMNS:
            if col not in df.columns:
                df[col] = None
        iid = db.create_import(conn, "frais-surgele", jour, f"np-{jour}.xlsx",
                               f"h-np-{jour}", "ok", nb_articles=len(rows))
        db.insert_snapshot(conn, iid, "frais-surgele", jour, df)


@pytest.fixture()
def seeded_negpos(fresh_db, monkeypatch):
    monkeypatch.setattr(config, "rayon_ids", lambda: ["frais-surgele", "epicerie-salee"])
    with fresh_db.lock_conn() as conn:
        _seed_negpos(conn)
    return fresh_db


def _run_negpos(tmp_path, monkeypatch, plan):
    monkeypatch.setattr(xi, "EXPORTS_DIR", str(tmp_path))
    monkeypatch.setattr(xi, "PUBLIC_BASE", "https://x/exports")
    monkeypatch.setattr(xi, "LAST_PLAN_FILE", str(tmp_path / "last.json"))
    return xi.build_excel(plan)


def test_marge_negative_stock_positif_meme_jour(seeded_negpos, tmp_path, monkeypatch):
    plan = {"rayon": "frais-surgele", "indicateur": "marge",
            "selection": {"marge_negative_stock_positif": True},
            "date_debut": "2026-07-30", "date_fin": "2026-08-01",
            "double_classement": True, "couleurs": True, "resume": True}
    res = _run_negpos(tmp_path, monkeypatch, plan)
    assert res["success"], res.get("erreur")
    assert res["nb_articles"] == 3, res
    assert res["partiel"] is False
    from openpyxl import load_workbook
    wb = load_workbook(res["fichier"])
    assert set(wb.sheetnames) == {"Prios argent", "Prix à corriger", "Résumé"}
    codes = {wb["Prios argent"].cell(row=r, column=1).value
             for r in range(2, wb["Prios argent"].max_row + 1)}
    assert codes == {900001, 900005, 900006}, codes
    f = res["filtre"]
    assert (f["nb_articles"], f["nb_journees"], f["premiere_occurrence"],
            f["derniere_occurrence"], f["marge_min"], f["stock_max"]) == \
           (3, 4, "2026-07-30", "2026-07-31", -9.0, 10)


def test_plan_from_args_mappe_filtre_negpos():
    plan = xi.plan_from_args({"rayon": "frais-surgele", "indicateur": "marge",
                              "marge_negative": True, "stock_positif": True,
                              "exclure_stock_zero": True})
    assert plan["selection"] == {"marge_negative_stock_positif": True}


def test_appel_plat_luna_aboutit(seeded, tmp_path, monkeypatch):
    # Bout en bout : les args plats de Luna produisent un fichier verifie.
    monkeypatch.setattr(xi, "EXPORTS_DIR", str(tmp_path))
    monkeypatch.setattr(xi, "PUBLIC_BASE", "https://x/exports")
    monkeypatch.setattr(xi, "LAST_PLAN_FILE", str(tmp_path / "last.json"))
    plan = {"rayon": "frais-surgele"}
    plan.update(xi.plan_from_args({"type": "marge", "rayon": "Frais surgelé",
                                   "format": "xlsx", "resume": True}))
    plan["rayon"] = config.resolve_rayon(plan["rayon"])
    res = xi.build_excel(plan)
    assert res["success"], res.get("erreur")
    assert res["nb_articles"] == 1
